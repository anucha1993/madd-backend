from openpyxl import load_workbook
from pathlib import Path

path = Path(r"D:\Programing\madd\ThepExcel-Thailand-Tambon.xlsx")
wb = load_workbook(path, read_only=True, data_only=False)
print("SHEET_NAMES:", repr(wb.sheetnames))
for ws in wb.worksheets:
    rows = ws.iter_rows(values_only=True)
    header = next(rows, ())
    samples = []
    for _ in range(5):
        row = next(rows, None)
        if row is None:
            break
        samples.append(row)
    # Count all rows reported by the worksheet, including the header row.
    print("\nSHEET:", repr(ws.title))
    print("HEADER:", repr(header))
    print("SAMPLE_ROWS:")
    for row in samples:
        print(repr(row))
    print("ROW_COUNT_INCLUDING_HEADER:", ws.max_row)
    print("DATA_ROW_COUNT_EXCLUDING_HEADER:", max(ws.max_row - 1, 0))
wb.close()
